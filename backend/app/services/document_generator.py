"""
Word document generator service.

Creates .docx documents with images arranged according to configuration.
Supports multiple alignment options (left, center, right) and layouts (vertical, inline).
"""

import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm
from docx.table import Table

from app.config import settings
from app.core.constants import PAGE_SIZES
from app.core.exceptions import DocumentGenerationError
from app.core.logging import logger
from app.models.requests import DocumentConfig
from app.services.image_processor import ProcessedImage


# Mapping for paragraph alignment
ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}

# Mapping for table alignment
TABLE_ALIGNMENT_MAP = {
    "left": WD_TABLE_ALIGNMENT.LEFT,
    "center": WD_TABLE_ALIGNMENT.CENTER,
    "right": WD_TABLE_ALIGNMENT.RIGHT,
}


@dataclass
class GenerationResult:
    """Result of document generation."""

    filename: str
    path: Path
    pages: int
    images_placed: int
    file_size_bytes: int

    @property
    def file_size_kb(self) -> float:
        """File size in KB."""
        return round(self.file_size_bytes / 1024, 2)


class DocumentGenerator:
    """Word document generator with images."""

    def __init__(self, output_dir: Path | None = None) -> None:
        """
        Initialize the generator.

        Args:
            output_dir: Output directory for documents
        """
        self.output_dir = output_dir or settings.get_output_path()

    def _get_page_dimensions(
        self,
        page_size: str,
        orientation: str,
    ) -> tuple[float, float]:
        """
        Get page dimensions in centimeters.

        Args:
            page_size: Page size (carta, a4, etc.)
            orientation: Orientation (vertical, horizontal)

        Returns:
            Tuple (width, height) in centimeters
        """
        dimensions = PAGE_SIZES.get(page_size, PAGE_SIZES["carta"])
        width, height = dimensions.width, dimensions.height

        if orientation == "horizontal":
            width, height = height, width

        return width, height

    def _calculate_images_per_row(
        self,
        config: DocumentConfig,
        page_width: float,
    ) -> int:
        """
        Calculate how many images fit per row.

        Args:
            config: Document configuration
            page_width: Page width in cm

        Returns:
            Number of images per row
        """
        if config.images_per_row != "auto":
            return int(config.images_per_row)

        # Available space = page width - margins
        available_width = page_width - (config.margins_cm * 2)

        # Images that fit considering spacing
        image_total_width = config.image_width_cm + config.spacing_cm
        images_per_row = int(available_width / image_total_width)

        # Minimum 1, maximum 6
        return max(1, min(images_per_row, 6))

    def _calculate_rows_needed(
        self,
        total_images: int,
        images_per_row: int,
    ) -> int:
        """
        Calculate the number of rows needed.

        Args:
            total_images: Total images
            images_per_row: Images per row

        Returns:
            Number of rows
        """
        return math.ceil(total_images / images_per_row)

    def _setup_document(
        self,
        config: DocumentConfig,
    ) -> tuple[Document, float, float]:
        """
        Set up a new document with specified dimensions.

        Args:
            config: Document configuration

        Returns:
            Tuple (Document, page_width, page_height)
        """
        doc = Document()

        page_width, page_height = self._get_page_dimensions(
            config.page_size,
            config.page_orientation,
        )

        # Configure section
        section = doc.sections[0]
        section.page_width = Cm(page_width)
        section.page_height = Cm(page_height)
        section.left_margin = Cm(config.margins_cm)
        section.right_margin = Cm(config.margins_cm)
        section.top_margin = Cm(config.margins_cm)
        section.bottom_margin = Cm(config.margins_cm)

        return doc, page_width, page_height

    def _create_table(
        self,
        doc: Document,
        rows: int,
        cols: int,
    ) -> Table:
        """
        Create a table to arrange images.

        Args:
            doc: Word document
            rows: Number of rows
            cols: Number of columns

        Returns:
            Created table
        """
        table = doc.add_table(rows=rows, cols=cols)

        # Remove table borders
        tbl = table._tbl
        tbl_pr = tbl.tblPr
        if tbl_pr is None:
            tbl_pr = parse_xml(f"<w:tblPr {nsdecls('w')}/>")
            tbl.insert(0, tbl_pr)

        # Configure invisible borders
        tbl_borders = parse_xml(
            f"<w:tblBorders {nsdecls('w')}>"
            '<w:top w:val="none"/>'
            '<w:left w:val="none"/>'
            '<w:bottom w:val="none"/>'
            '<w:right w:val="none"/>'
            '<w:insideH w:val="none"/>'
            '<w:insideV w:val="none"/>'
            "</w:tblBorders>"
        )
        tbl_pr.append(tbl_borders)

        return table

    def _insert_image(
        self,
        cell,
        image: ProcessedImage,
        width_cm: float,
        height_cm: float | None,
        add_border: bool,
        alignment: str = "left",
    ) -> None:
        """
        Insert an image into a table cell.

        Args:
            cell: Table cell
            image: Processed image
            width_cm: Width in centimeters
            height_cm: Height in centimeters (None = proportional)
            add_border: Add border to image
            alignment: Horizontal alignment (left, center, right)
        """
        # Configure cell
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        # Get or create paragraph
        paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        paragraph.alignment = ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)

        # Reset image stream
        image.stream.seek(0)

        # Create run and insert image
        run = paragraph.add_run()

        # Calculate dimensions
        if height_cm is None:
            # Maintain proportion
            aspect_ratio = image.info.aspect_ratio
            height_cm = width_cm / aspect_ratio

        # Insert image with dimensions
        picture = run.add_picture(image.stream, width=Cm(width_cm), height=Cm(height_cm))

        # Add border if configured
        if add_border:
            self._add_picture_border(picture)

    def _add_picture_border(self, picture) -> None:
        """
        Add a black border to an image.

        Args:
            picture: InlineShape object of the image
        """
        # Get image XML element
        inline = picture._inline
        spPr = inline.graphic.graphicData.pic.spPr

        # Create line element
        ln = parse_xml(
            f'<a:ln {nsdecls("a")} w="12700">'  # 1pt = 12700 EMUs
            '<a:solidFill><a:srgbClr val="000000"/></a:solidFill>'
            "</a:ln>"
        )
        spPr.append(ln)

    def _insert_image_to_paragraph(
        self,
        paragraph,
        image: ProcessedImage,
        width_cm: float,
        height_cm: float | None,
        add_border: bool,
    ) -> None:
        """
        Insert an image into a paragraph.

        Args:
            paragraph: Document paragraph
            image: Processed image
            width_cm: Width in centimeters
            height_cm: Height in centimeters (None = proportional)
            add_border: Add border to image
        """
        # Reset image stream
        image.stream.seek(0)

        # Create run and insert image
        run = paragraph.add_run()

        # Calculate dimensions
        if height_cm is None:
            aspect_ratio = image.info.aspect_ratio
            height_cm = width_cm / aspect_ratio

        # Insert image with dimensions
        picture = run.add_picture(image.stream, width=Cm(width_cm), height=Cm(height_cm))

        # Add border if configured
        if add_border:
            self._add_picture_border(picture)

    def _generate_vertical_layout(
        self,
        doc: Document,
        images: list[ProcessedImage],
        config: DocumentConfig,
    ) -> None:
        """
        Generate document with vertical layout (one image per line).

        Args:
            doc: Word document
            images: List of processed images
            config: Document configuration
        """
        alignment = ALIGNMENT_MAP.get(config.image_alignment, WD_ALIGN_PARAGRAPH.LEFT)

        for image in images:
            paragraph = doc.add_paragraph()
            paragraph.alignment = alignment

            self._insert_image_to_paragraph(
                paragraph=paragraph,
                image=image,
                width_cm=config.image_width_cm,
                height_cm=config.image_height_cm,
                add_border=config.borders,
            )

    def _generate_inline_layout(
        self,
        doc: Document,
        images: list[ProcessedImage],
        config: DocumentConfig,
        page_width: float,
    ) -> None:
        """
        Generate document with inline layout (multiple images per line).

        Args:
            doc: Word document
            images: List of processed images
            config: Document configuration
            page_width: Page width in cm
        """
        images_per_row = self._calculate_images_per_row(config, page_width)
        rows_needed = self._calculate_rows_needed(len(images), images_per_row)

        # Create table for inline layout
        table = self._create_table(doc, rows_needed, images_per_row)

        # Set table alignment
        table_alignment = TABLE_ALIGNMENT_MAP.get(config.image_alignment, WD_TABLE_ALIGNMENT.LEFT)
        table.alignment = table_alignment

        # Insert images
        for idx, image in enumerate(images):
            row = idx // images_per_row
            col = idx % images_per_row

            cell = table.rows[row].cells[col]
            self._insert_image(
                cell=cell,
                image=image,
                width_cm=config.image_width_cm,
                height_cm=config.image_height_cm,
                add_border=config.borders,
                alignment=config.image_alignment,
            )

    def _estimate_pages(
        self,
        total_images: int,
        images_per_row: int,
        image_height_cm: float,
        page_height: float,
        margins_cm: float,
        spacing_cm: float,
    ) -> int:
        """
        Estimate the number of pages in the document.

        Args:
            total_images: Total images
            images_per_row: Images per row
            image_height_cm: Height of each image
            page_height: Page height
            margins_cm: Margins
            spacing_cm: Space between images

        Returns:
            Estimated number of pages
        """
        # Available space per page
        available_height = page_height - (margins_cm * 2)

        # Height of each row (image + spacing)
        row_height = image_height_cm + spacing_cm

        # Rows per page
        rows_per_page = max(1, int(available_height / row_height))

        # Total rows needed
        total_rows = math.ceil(total_images / images_per_row)

        # Pages needed
        return max(1, math.ceil(total_rows / rows_per_page))

    def generate(
        self,
        images: list[ProcessedImage],
        config: DocumentConfig,
    ) -> GenerationResult:
        """
        Generate a Word document with images.

        Args:
            images: List of processed images
            config: Document configuration

        Returns:
            GenerationResult with information about the generated document

        Raises:
            DocumentGenerationError: If there's an error in generation
        """
        if not images:
            raise DocumentGenerationError("No images to generate document")

        try:
            logger.info(
                f"Generating document: {len(images)} images, "
                f"size={config.page_size}, image_width={config.image_width_cm}cm, "
                f"alignment={config.image_alignment}, layout={config.image_layout}"
            )

            # 1. Set up document
            doc, page_width, page_height = self._setup_document(config)

            # 2. Generate based on layout type
            if config.image_layout == "vertical":
                # Vertical layout: one image per line
                self._generate_vertical_layout(doc, images, config)
                images_per_row = 1
            else:
                # Inline layout: multiple images per row in a table
                self._generate_inline_layout(doc, images, config, page_width)
                images_per_row = self._calculate_images_per_row(config, page_width)

            rows_needed = self._calculate_rows_needed(len(images), images_per_row)

            logger.debug(
                f"Layout: {config.image_layout}, {images_per_row} images/row, "
                f"{rows_needed} total rows, alignment={config.image_alignment}"
            )

            # 3. Generate filename
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"{config.filename}_{timestamp}.docx"
            output_path = self.output_dir / filename

            # 4. Save document
            doc.save(str(output_path))

            # 5. Get file size
            file_size = output_path.stat().st_size

            # 6. Estimate pages
            avg_image_height = config.image_width_cm / 1.5  # Typical aspect ratio
            estimated_pages = self._estimate_pages(
                len(images),
                images_per_row,
                config.image_height_cm or avg_image_height,
                page_height,
                config.margins_cm,
                config.spacing_cm,
            )

            result = GenerationResult(
                filename=filename,
                path=output_path,
                pages=estimated_pages,
                images_placed=len(images),
                file_size_bytes=file_size,
            )

            logger.info(
                f"Document generated: {filename} "
                f"({result.file_size_kb}KB, ~{estimated_pages} pages)"
            )

            return result

        except DocumentGenerationError:
            raise
        except Exception as e:
            raise DocumentGenerationError(
                f"Error generating document: {e}",
                detail=str(type(e).__name__),
            ) from e


# Singleton instance
document_generator = DocumentGenerator()
